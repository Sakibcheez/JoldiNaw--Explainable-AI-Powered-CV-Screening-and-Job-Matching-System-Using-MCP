import io
import re
import zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfReader

SKILLS = [
    "Python",
    "FastAPI",
    "React",
    "TypeScript",
    "JavaScript",
    "SQL",
    "NoSQL",
    "PostgreSQL",
    "Docker",
    "Redis",
    "Git",
    "C",
    "C++",
    "C#",
    "Node.js",
    ".NET",
    "Java",
    "AWS",
    "Kubernetes",
    "Machine Learning",
]


def has_term(text: str, term: str):
    return bool(re.search(r"(?<![\w+#.])" + re.escape(term) + r"(?![\w+#])", text, re.I))


def skills(text: str):
    return [skill for skill in SKILLS if has_term(text, skill)]


def validate_file(data: bytes, filename: str):
    suffix = Path(filename).suffix.lower()
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("File exceeds 10 MB")
    if suffix == ".pdf" and data.startswith(b"%PDF-"):
        return suffix
    if suffix == ".docx" and data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                    raise ValueError("Expanded document exceeds limits")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("Invalid DOCX")
                if any("vbaProject" in e.filename or e.flag_bits & 1 for e in entries):
                    raise ValueError("Encrypted or macro-enabled document rejected")
            return suffix
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid DOCX") from exc
    raise ValueError("Upload a valid PDF or DOCX")


def passages(text: str, page: int = 1):
    # Bounded passages retain offsets in the extracted page, including whitespace.
    for start in range(0, len(text), 900):
        chunk = text[start : start + 900]
        if chunk.strip():
            yield {"text": chunk, "page": page, "start": start, "end": start + len(chunk)}


def parse(data: bytes, filename: str):
    suffix = validate_file(data, filename)
    output = []
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are unsupported")
        if len(reader.pages) > 50:
            raise ValueError("Document exceeds 50 pages")
        for i, page in enumerate(reader.pages, 1):
            output.extend(passages(page.extract_text() or "", i))
    else:
        doc = Document(io.BytesIO(data))
        text = "\n".join(
            [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
        )
        output.extend(passages(text))
    if not output:
        raise ValueError("No extractable text; scanned documents require OCR")
    if sum(len(p["text"]) for p in output) > 100000:
        raise ValueError("Extracted text exceeds limits")
    return output
